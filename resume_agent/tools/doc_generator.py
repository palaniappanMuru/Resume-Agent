from pathlib import Path

from docx import Document
from docx.shared import Pt
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import HumanMessage, SystemMessage

from resume_agent.config import settings
from resume_agent.schemas import (
    ATSReport,
    GraphContext,
    JDMatchStatus,
    JDRequirements,
    ResumeContent,
    TailoredResume,
)

CV_SYSTEM_PROMPT = """You are an expert resume writer. Using ONLY the candidate evidence provided \
below (matched skills, projects, accomplishments, courses), write a tailored resume for the given \
job title. Rules:
- summary: 2-3 sentences, professional tone, tailored to the job title and seniority level given.
- skills: the candidate's matched skills, ordered with the most relevant to the job first.
- experience: one entry per project provided, using its accomplishments as bullets. Bullets may be \
rephrased for clarity/impact but must not invent facts, numbers, or technologies not present in the \
evidence.Just provide the job roles only without description for the roles before 2013.Mention the duration of the projects.
- education: one entry per course provided.
- Max size: 2 pages . Use concise language and avoid filler words.
Do not include any project, skill, or course not present in the evidence. Do not fabricate evidence."""


def _evidence_block(jd: JDRequirements, graph_ctx: GraphContext, role_title: str) -> str:
    matched_skill_names = sorted({ms.graph_skill for ms in graph_ctx.matched_skills})
    lines = [
        f"Job title: {role_title}",
        f"Seniority level: {jd.seniority_level or 'unspecified'}",
        f"Matched skills: {', '.join(matched_skill_names) or 'none'}",
        "Projects:",
    ]
    for p in graph_ctx.projects:
        lines.append(f"- {p.name}: {p.description or ''}")
        for bullet in p.accomplishments:
            lines.append(f"  - {bullet}")
    lines.append("Courses:")
    for c in graph_ctx.courses:
        lines.append(f"- {c.name} (skill: {c.skill or 'n/a'})")
    return "\n".join(lines)


def build_resume(
    jd: JDRequirements,
    graph_ctx: GraphContext,
    ats_report: ATSReport,
    jd_match_status: JDMatchStatus,
    llm: BaseChatModel,
    job_role_name: str | None = None,
) -> TailoredResume:
    """Compile JD + retrieved graph context + ATS report into the final tailored resume data.

    The resume's summary/skills/experience/education are authored by `llm` (the configurable
    CV-construction model), grounded strictly in the evidence retrieved from the graph.
    `job_role_name` is the deduped JobRole name resolved by graph_writer.sync_job_role_and_skills
    (may differ in casing/wording from jd.job_title if an existing role was reused); falls back
    to jd.job_title when not provided.
    """
    role_title = job_role_name or jd.job_title

    structured_llm = llm.with_structured_output(ResumeContent)
    content: ResumeContent = structured_llm.invoke(
        [
            SystemMessage(content=CV_SYSTEM_PROMPT),
            HumanMessage(content=_evidence_block(jd, graph_ctx, role_title)),
        ]
    )

    return TailoredResume(
        candidate_name=settings.candidate_name or "Candidate Name",
        contact={
            "email": settings.candidate_email,
            "phone": settings.candidate_phone,
            "location": settings.candidate_location,
            "linkedin": settings.candidate_linkedin,
        },
        job_title=role_title,
        summary=content.summary,
        skills=content.skills,
        experience=content.experience,
        education=content.education,
        ats_report=ats_report,
        jd_match_status=jd_match_status,
    )


def _slug(text: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in text.lower()).strip("_") or "resume"


def _write_docx(resume: TailoredResume, path: Path) -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(resume.candidate_name, level=0)
    contact_line = " | ".join(v for v in resume.contact.values() if v)
    if contact_line:
        doc.add_paragraph(contact_line)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(resume.summary)

    if resume.skills:
        doc.add_heading("Core Skills", level=1)
        doc.add_paragraph(", ".join(resume.skills))

    if resume.experience:
        doc.add_heading("Experience", level=1)
        for entry in resume.experience:
            doc.add_heading(entry.project, level=2)
            for bullet in entry.bullets:
                doc.add_paragraph(bullet, style="List Bullet")

    if resume.education:
        doc.add_heading("Education & Certifications", level=1)
        for entry in resume.education:
            line = entry.course
            if entry.related_skill:
                line += f" ({entry.related_skill})"
            doc.add_paragraph(line, style="List Bullet")

    doc.save(str(path))


def generate_documents(resume: TailoredResume, output_dir: str | Path | None = None) -> tuple[Path, Path]:
    """Export the tailored resume as both machine-readable JSON and an ATS-friendly .docx.

    Task 2.4: Document Generation Tool.
    Returns (json_path, docx_path).
    """
    out_dir = Path(output_dir or settings.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    slug = _slug(resume.job_title)
    json_path = out_dir / f"{slug}_resume.json"
    docx_path = out_dir / f"{slug}_resume.docx"

    json_path.write_text(resume.model_dump_json(indent=2), encoding="utf-8")
    _write_docx(resume, docx_path)

    return json_path, docx_path
